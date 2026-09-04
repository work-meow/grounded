"""The parser is dispatched on magic bytes and feeds page numbers to citations.

Every format is built here and read back, so a library upgrade that changes the
shape of what it returns fails loudly instead of quietly producing empty chunks.
"""

import io

import pytest
from docx import Document
from fpdf import FPDF
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from rag_indexer.parsers import parse_document


@pytest.fixture(scope="module")
def pdf_bytes() -> bytes:
    pdf = FPDF()
    for page in range(1, 4):
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        pdf.multi_cell(0, 8, f"Page {page} says notice is {page * 10} days.")
    return bytes(pdf.output())


def test_pdf_yields_one_entry_per_page_with_its_number(pdf_bytes):
    parsed = parse_document(pdf_bytes)

    assert [meta["page_number"] for _, meta in parsed] == [1, 2, 3]
    # The page number must belong to the text it is attached to, or citations
    # would point at the wrong page.
    assert "notice is 20 days" in parsed[1][0]


def test_docx_reads_paragraphs_and_tables():
    document = Document()
    document.add_paragraph("Договор аренды помещения.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Срок уведомления"
    table.rows[0].cells[1].text = "30 дней"
    buffer = io.BytesIO()
    document.save(buffer)

    ((text, meta),) = parse_document(buffer.getvalue())

    assert "Договор аренды" in text
    # Contracts keep their terms in tables at least as often as in prose.
    assert "Срок уведомления | 30 дней" in text
    assert meta == {}


def test_pptx_numbers_slides_like_pages():
    presentation = Presentation()
    for slide_text in ("Первый слайд", "Второй слайд"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        box.text_frame.text = slide_text
    buffer = io.BytesIO()
    presentation.save(buffer)

    parsed = parse_document(buffer.getvalue())

    assert [meta["page_number"] for _, meta in parsed] == [1, 2]
    assert "Второй слайд" in parsed[1][0]


def test_xlsx_flattens_rows_per_sheet():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Условия"
    sheet.append(["Пункт", "Значение"])
    sheet.append(["Уведомление", 30])
    buffer = io.BytesIO()
    workbook.save(buffer)

    ((text, meta),) = parse_document(buffer.getvalue())

    assert "Уведомление | 30" in text
    assert meta["sheet"] == "Условия"


@pytest.mark.parametrize(
    "raw",
    [
        "Обычный текст договора.".encode(),
        "Кириллица в cp1251.".encode("cp1251"),
        b"# Markdown\n\nWith **bold** text.\n",
    ],
)
def test_plain_text_survives_several_encodings(raw):
    ((text, _),) = parse_document(raw)
    assert text.strip()


def test_unreadable_input_returns_nothing_instead_of_raising():
    """One bad upload must not take the whole indexing pipeline down."""
    assert parse_document(b"%PDF-1.4 but truncated and broken") == []
    assert parse_document(b"") == []


def test_zip_that_is_not_an_office_document_is_ignored():
    buffer = io.BytesIO()
    with __import__("zipfile").ZipFile(buffer, "w") as archive:
        archive.writestr("random.txt", "not office")

    assert parse_document(buffer.getvalue()) == []
