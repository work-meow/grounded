"""Document parsing, dispatched on the file's own magic bytes.

Pathway hands a parser raw bytes and no filename, so the format is sniffed from
the content — which is more trustworthy than an extension in any case.

This replaces ``unstructured``, which could not do the job here:

* PDF required the ``[pdf]`` extra even for ``strategy="fast"``. The dependency
  check in its partitioner loader runs before the strategy is consulted, so PDFs
  raised ImportError; satisfying it means unstructured-inference, torch and CUDA
  — about 2 GB — for layout models that a fast text extraction never touches.
* DOCX took 7.6 s for 30 paragraphs, against 34 ms for python-docx, because
  every element goes through NLTK.
* It weighed ~300 MB plus NLTK corpora plus libmagic, against ~14 MB for the
  four libraries used here.

Each parser returns ``(text, metadata)`` pairs. Paged formats return one pair
per page so that ``page_number`` survives into every chunk and a citation can
say "contract.pdf, стр. 14".
"""

import io
import logging
import zipfile

import pypdfium2 as pdfium

_PDF_MAGIC = b"%PDF-"
_ZIP_MAGIC = b"PK\x03\x04"

logger = logging.getLogger(__name__)


def parse_document(contents: bytes) -> list[tuple[str, dict]]:
    """Turn a file into ``(text, metadata)`` pairs.

    Returns an empty list for anything unreadable rather than raising: one bad
    upload must not take the indexing pipeline down with it. The document then
    stays visibly at "processing" in the UI, which is the honest signal.
    """
    try:
        if contents.startswith(_PDF_MAGIC):
            return _pdf(contents)
        if contents.startswith(_ZIP_MAGIC):
            return _ooxml(contents)
        return _plain_text(contents)
    except Exception:
        logger.exception("could not parse a %d-byte document", len(contents))
        return []


def _pdf(contents: bytes) -> list[tuple[str, dict]]:
    """One entry per page, via PDFium.

    ~5x faster than the alternatives measured and the only one of them that
    reports page boundaries, which the citations depend on.
    """
    pages: list[tuple[str, dict]] = []
    document = pdfium.PdfDocument(contents)
    try:
        for number, page in enumerate(document, start=1):
            textpage = page.get_textpage()
            try:
                text = textpage.get_text_bounded()
            finally:
                # These wrap C++ handles; letting the GC decide would hold
                # native memory for as long as it feels like.
                textpage.close()
                page.close()
            if text.strip():
                pages.append((text, {"page_number": number}))
    finally:
        document.close()
    return pages


def _ooxml(contents: bytes) -> list[tuple[str, dict]]:
    """DOCX, PPTX and XLSX are all ZIPs; the part names tell them apart."""
    with zipfile.ZipFile(io.BytesIO(contents)) as archive:
        names = set(archive.namelist())

    if "word/document.xml" in names:
        return _docx(contents)
    if "ppt/presentation.xml" in names:
        return _pptx(contents)
    if "xl/workbook.xml" in names:
        return _xlsx(contents)

    logger.warning("zip archive is not a recognised Office document")
    return []


def _docx(contents: bytes) -> list[tuple[str, dict]]:
    """Whole document as one entry — Word has no page breaks until it is laid out."""
    from docx import Document

    document = Document(io.BytesIO(contents))
    blocks = [p.text for p in document.paragraphs if p.text.strip()]
    # Contracts keep their terms in tables far more often than in prose.
    blocks += [
        " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
        for table in document.tables
        for row in table.rows
        if any(cell.text.strip() for cell in row.cells)
    ]
    text = "\n".join(blocks)
    return [(text, {})] if text.strip() else []


def _pptx(contents: bytes) -> list[tuple[str, dict]]:
    """One entry per slide; the slide number doubles as the page number."""
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(contents))
    slides: list[tuple[str, dict]] = []
    for number, slide in enumerate(presentation.slides, start=1):
        text = "\n".join(
            shape.text_frame.text
            for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
        )
        if text.strip():
            slides.append((text, {"page_number": number}))
    return slides


def _xlsx(contents: bytes) -> list[tuple[str, dict]]:
    """One entry per sheet, rows flattened.

    read_only keeps openpyxl from building the whole object model, which is
    what makes a large workbook expensive.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
    try:
        sheets: list[tuple[str, dict]] = []
        for sheet in workbook.worksheets:
            rows = [
                " | ".join(str(cell) for cell in row if cell is not None)
                for row in sheet.iter_rows(values_only=True)
            ]
            text = "\n".join(row for row in rows if row.strip())
            if text.strip():
                sheets.append((text, {"sheet": sheet.title}))
        return sheets
    finally:
        workbook.close()


def _plain_text(contents: bytes) -> list[tuple[str, dict]]:
    """TXT and Markdown. Latin-1 never fails, so it is the last resort."""
    for encoding in ("utf-8", "utf-16", "cp1251", "latin-1"):
        try:
            text = contents.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    return [(text, {})] if text.strip() else []
